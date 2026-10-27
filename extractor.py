#!/usr/bin/env python3
"""
Universal embedded-document extractor and text reader.

Set INPUT_PATH in CONFIGURATION. Run this file without command-line arguments.

Supported parent files:
- PDF: embedded-file collection, file-attachment annotations, and embedded images
- DOCX/XLSX/PPTX and other OOXML/ZIP containers: embedded package files,
  OLE package objects, and document images
- Other files: text is extracted from the file itself when supported

Supported text extraction:
PDF, DOCX, XLSX/XLSM, PPTX, CSV/TSV, TXT/MD/LOG, JSON/XML/HTML,
RTF, common images through OCR, and legacy Office files through LibreOffice
when it is installed.
"""

from __future__ import annotations

import base64
import csv
import hashlib
import io
import json
import mimetypes
import re
import shutil
import struct
import subprocess
import urllib.error
import urllib.parse
import urllib.request
import zipfile
from dataclasses import asdict, dataclass
from html.parser import HTMLParser
from pathlib import Path
from typing import Any
import xml.etree.ElementTree as ET


# ============================================================================
# CONFIGURATION: EDIT THE PATH HERE. No terminal argument is required.
# ============================================================================
INPUT_PATH = r"/home/mridulj/process_compiler_v2/input/BRD on AI-driven Design Validation.docx"

# Output is created beside the input file unless an absolute path is supplied.
OUTPUT_FOLDER_NAME = "/home/mridulj/process_compiler_v2/link_testing/extracted_documents_brd"

# Embedded documents, images, and external document links are processed.
MAX_RECURSION_DEPTH = 5
DOWNLOAD_EXTERNAL_LINKS = True
DOWNLOAD_TIMEOUT_SECONDS = 45
MAX_DOWNLOAD_BYTES = 200 * 1024 * 1024

# Optional. Leave blank for public/anonymous SharePoint links.
# For protected SharePoint links, paste a valid Microsoft Graph access token here.
# Do not commit a real token to source control.
SHAREPOINT_ACCESS_TOKEN = ""

# Print all extracted text after processing.
PRINT_CONTENT_TO_CONSOLE = True


@dataclass
class ExtractedItem:
    source: str
    saved_path: str | None
    kind: str
    content_type: str | None = None
    text: str = ""
    error: str | None = None
    link: str | None = None
    access_status: str | None = None


# Public result variables populated by main().
EXTRACTED_CONTENTS: dict[str, str] = {}
ALL_CONTENT_TEXT: str = ""
RESULT: dict[str, Any] = {}


SIGNATURE_EXTENSIONS = [
    (b"%PDF-", ".pdf"),
    (b"PK\x03\x04", ".zip"),
    (b"\xD0\xCF\x11\xE0\xA1\xB1\x1A\xE1", ".ole"),
    (b"\x89PNG\r\n\x1a\n", ".png"),
    (b"\xff\xd8\xff", ".jpg"),
    (b"GIF87a", ".gif"),
    (b"GIF89a", ".gif"),
    (b"BM", ".bmp"),
    (b"II*\x00", ".tif"),
    (b"MM\x00*", ".tif"),
    (b"RIFF", ".webp"),
]

OOXML_MAIN_TYPES = {
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.template": ".dotx",
    "application/vnd.ms-word.document.macroEnabled.12": ".docm",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": ".xlsx",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.template": ".xltx",
    "application/vnd.ms-excel.sheet.macroEnabled.12": ".xlsm",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": ".pptx",
    "application/vnd.openxmlformats-officedocument.presentationml.template": ".potx",
    "application/vnd.openxmlformats-officedocument.presentationml.slideshow": ".ppsx",
    "application/vnd.ms-powerpoint.presentation.macroEnabled.12": ".pptm",
}

GENERIC_EXTENSIONS = {"", ".bin", ".dat", ".ole", ".tmp", ".zip"}


def safe_name(name: str, fallback: str = "embedded_file") -> str:
    name = urllib.parse.unquote(name or "")
    name = Path(name.replace("\\", "/")).name.strip().strip(".")
    name = re.sub(r"[<>:\"/\\|?*\x00-\x1f]", "_", name)
    return name[:180] or fallback


def unique_path(folder: Path, filename: str) -> Path:
    folder.mkdir(parents=True, exist_ok=True)
    filename = safe_name(filename)
    candidate = folder / filename
    if not candidate.exists():
        return candidate
    stem, suffix = candidate.stem, candidate.suffix
    i = 2
    while True:
        candidate = folder / f"{stem}_{i}{suffix}"
        if not candidate.exists():
            return candidate
        i += 1


def read_head(path: Path, n: int = 8192) -> bytes:
    with path.open("rb") as f:
        return f.read(n)


def detect_ooxml_extension(path: Path) -> str | None:
    try:
        if not zipfile.is_zipfile(path):
            return None
        with zipfile.ZipFile(path) as zf:
            if "[Content_Types].xml" not in zf.namelist():
                return None
            root = ET.fromstring(zf.read("[Content_Types].xml"))
            # Only inspect the package main part. Embedded Office files can add
            # other Office content types to [Content_Types].xml.
            main_parts = (
                "/word/document.xml",
                "/xl/workbook.xml",
                "/ppt/presentation.xml",
            )
            for elem in root:
                part_name = elem.attrib.get("PartName", "")
                ctype = elem.attrib.get("ContentType", "")
                if part_name == "/word/document.xml":
                    return ".docm" if "macroEnabled" in ctype else ".docx"
                if part_name == "/xl/workbook.xml":
                    return ".xlsm" if "macroEnabled" in ctype else ".xlsx"
                if part_name == "/ppt/presentation.xml":
                    return ".pptm" if "macroEnabled" in ctype else ".pptx"
    except Exception:
        return None
    return None


def detect_extension(path: Path, content_type: str | None = None,
                     suggested_name: str | None = None) -> str:
    head = read_head(path)
    ooxml_ext = detect_ooxml_extension(path)
    if ooxml_ext:
        return ooxml_ext

    # Handle a few signature-specific formats first.
    if head.startswith(b"{\\rtf"):
        return ".rtf"
    if head.startswith(b"RIFF") and head[8:12] == b"WEBP":
        return ".webp"

    for signature, ext in SIGNATURE_EXTENSIONS:
        if head.startswith(signature):
            return ext

    if content_type:
        clean_type = content_type.split(";", 1)[0].strip().lower()
        if clean_type in OOXML_MAIN_TYPES:
            return OOXML_MAIN_TYPES[clean_type]
        guessed = mimetypes.guess_extension(clean_type)
        if guessed:
            return ".jpg" if guessed == ".jpe" else guessed

    suffix = Path(suggested_name or path.name).suffix.lower()
    return suffix or ".bin"


def ensure_correct_extension(path: Path, content_type: str | None = None,
                             suggested_name: str | None = None) -> Path:
    detected = detect_extension(path, content_type, suggested_name)
    current = path.suffix.lower()
    # Preserve a meaningful supplied extension, except when signature/OOXML proves otherwise.
    ooxml_ext = detect_ooxml_extension(path)
    head = read_head(path, 32)
    signature_ext = None
    if ooxml_ext:
        signature_ext = ooxml_ext
    else:
        for signature, ext in SIGNATURE_EXTENSIONS:
            if head.startswith(signature):
                signature_ext = ext
                break
        if head.startswith(b"{\\rtf"):
            signature_ext = ".rtf"
        if head.startswith(b"RIFF") and head[8:12] == b"WEBP":
            signature_ext = ".webp"

    desired = signature_ext or (detected if current in GENERIC_EXTENSIONS else current)
    if desired == ".zip" and current not in GENERIC_EXTENSIONS:
        desired = current
    if desired and current != desired:
        new_path = unique_path(path.parent, path.stem + desired)
        path.replace(new_path)
        return new_path
    return path


def write_bytes_validated(folder: Path, name: str, data: bytes,
                          content_type: str | None = None) -> Path:
    target = unique_path(folder, safe_name(name))
    target.write_bytes(data)
    return ensure_correct_extension(target, content_type, name)


def file_hash(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16", "utf-16-le", "utf-16-be", "cp1252", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            pass
    return data.decode("utf-8", errors="replace")


class _HTMLTextExtractor(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []

    def handle_data(self, data: str) -> None:
        if data.strip():
            self.parts.append(data.strip())


def html_to_text(raw: str) -> str:
    parser = _HTMLTextExtractor()
    parser.feed(raw)
    return "\n".join(parser.parts)


def extract_pdf_text(path: Path) -> str:
    import pymupdf
    doc = pymupdf.open(path)
    pages = []
    for i, page in enumerate(doc, 1):
        text = page.get_text("text", sort=True).strip()
        if not text:
            # OCR only if Tesseract support is available in the environment.
            try:
                tp = page.get_textpage_ocr(dpi=300, full=True)
                text = page.get_text("text", textpage=tp, sort=True).strip()
            except Exception:
                pass
        pages.append(f"\n--- Page {i} ---\n{text}")
    doc.close()
    return "".join(pages).strip()


def extract_docx_text(path: Path) -> str:
    from docx import Document
    doc = Document(path)
    parts: list[str] = []
    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            parts.append(paragraph.text)
    for table_no, table in enumerate(doc.tables, 1):
        parts.append(f"\n[Table {table_no}]")
        for row in table.rows:
            parts.append("\t".join(cell.text.replace("\n", " ").strip() for cell in row.cells))
    for section in doc.sections:
        for paragraph in section.header.paragraphs:
            if paragraph.text.strip():
                parts.append("[Header] " + paragraph.text)
        for paragraph in section.footer.paragraphs:
            if paragraph.text.strip():
                parts.append("[Footer] " + paragraph.text)
    return "\n".join(parts)


def extract_xlsx_text(path: Path) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
    parts: list[str] = []
    try:
        for ws in wb.worksheets:
            parts.append(f"\n=== Sheet: {ws.title} ===")
            for row in ws.iter_rows(values_only=True):
                values = ["" if v is None else str(v) for v in row]
                # Keep internal empty cells, trim only trailing empty cells.
                while values and values[-1] == "":
                    values.pop()
                if values:
                    parts.append("\t".join(values))
    finally:
        wb.close()
    return "\n".join(parts).strip()


def extract_pptx_text(path: Path) -> str:
    from pptx import Presentation
    prs = Presentation(path)
    parts: list[str] = []
    for slide_no, slide in enumerate(prs.slides, 1):
        parts.append(f"\n=== Slide {slide_no} ===")
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                parts.append(shape.text)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    parts.append("\t".join(cell.text.strip() for cell in row.cells))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append("[Notes] " + notes)
    return "\n".join(parts).strip()


def extract_csv_text(path: Path) -> str:
    raw = path.read_bytes()
    text = decode_text(raw)
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",\t;|")
    except csv.Error:
        dialect = csv.excel_tab if path.suffix.lower() == ".tsv" else csv.excel
    rows = csv.reader(io.StringIO(text), dialect)
    return "\n".join("\t".join(row) for row in rows)


def extract_rtf_text(path: Path) -> str:
    try:
        from striprtf.striprtf import rtf_to_text
        return rtf_to_text(decode_text(path.read_bytes()))
    except ImportError:
        # Basic fallback. striprtf gives better results.
        raw = decode_text(path.read_bytes())
        raw = re.sub(r"\\par[d]?\b", "\n", raw)
        raw = re.sub(r"\\'[0-9a-fA-F]{2}", "", raw)
        raw = re.sub(r"\\[a-zA-Z]+-?\d* ?", "", raw)
        return re.sub(r"[{}]", "", raw)


def extract_image_text(path: Path) -> str:
    # RapidOCR does not need a separately installed Tesseract executable.
    try:
        from rapidocr_onnxruntime import RapidOCR
        engine = RapidOCR()
        result, _ = engine(str(path))
        if result:
            return "\n".join(str(item[1]) for item in result if len(item) > 1)
    except ImportError:
        pass
    except Exception as exc:
        rapid_error = str(exc)
    else:
        rapid_error = "no text detected"

    # Optional fallback when Tesseract is already available.
    try:
        import pytesseract
        from PIL import Image
        return pytesseract.image_to_string(Image.open(path)).strip()
    except Exception as exc:
        raise RuntimeError(
            "Image OCR unavailable. Install rapidocr-onnxruntime, or install "
            f"Tesseract plus pytesseract. RapidOCR status: {locals().get('rapid_error', 'not installed')}; "
            f"Tesseract status: {exc}"
        )


def libreoffice_convert(path: Path, target_ext: str, temp_dir: Path) -> Path | None:
    executable = shutil.which("libreoffice") or shutil.which("soffice")
    if not executable:
        return None
    temp_dir.mkdir(parents=True, exist_ok=True)
    fmt = {".docx": "docx", ".xlsx": "xlsx", ".pptx": "pptx"}[target_ext]
    cmd = [executable, "--headless", "--convert-to", fmt, "--outdir", str(temp_dir), str(path)]
    completed = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, timeout=180)
    candidates = list(temp_dir.glob(path.stem + ".*"))
    expected = [p for p in candidates if p.suffix.lower() == target_ext]
    return expected[0] if completed.returncode == 0 and expected else None


def extract_text(path: Path, work_dir: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return extract_pdf_text(path)
    if suffix in {".docx", ".docm", ".dotx"}:
        return extract_docx_text(path)
    if suffix in {".xlsx", ".xlsm", ".xltx"}:
        return extract_xlsx_text(path)
    if suffix in {".pptx", ".pptm", ".ppsx", ".potx"}:
        return extract_pptx_text(path)
    if suffix in {".csv", ".tsv"}:
        return extract_csv_text(path)
    if suffix in {".txt", ".md", ".log", ".ini", ".cfg", ".yaml", ".yml", ".sql", ".py"}:
        return decode_text(path.read_bytes())
    if suffix == ".json":
        obj = json.loads(decode_text(path.read_bytes()))
        return json.dumps(obj, indent=2, ensure_ascii=False, default=str)
    if suffix in {".xml", ".html", ".htm"}:
        raw = decode_text(path.read_bytes())
        return html_to_text(raw) if suffix in {".html", ".htm"} else raw
    if suffix == ".rtf":
        return extract_rtf_text(path)
    if suffix in {".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff", ".gif", ".webp"}:
        return extract_image_text(path)
    if suffix in {".doc", ".xls", ".ppt"}:
        target = {".doc": ".docx", ".xls": ".xlsx", ".ppt": ".pptx"}[suffix]
        converted = libreoffice_convert(path, target, work_dir / "converted")
        if converted:
            return extract_text(converted, work_dir)
        raise RuntimeError("Legacy Office file requires LibreOffice/soffice to be installed.")
    return ""


def parse_ole10native(data: bytes) -> tuple[str, bytes] | None:
    """Parse the common OLE Package stream used by Office embedded objects."""
    try:
        pos = 4  # total native data size
        pos += 2 # flags
        end = data.index(b"\x00", pos)
        label = decode_text(data[pos:end]).strip()
        pos = end + 1
        end = data.index(b"\x00", pos)
        _original_path = data[pos:end]
        pos = end + 1
        pos += 8  # reserved fields
        end = data.index(b"\x00", pos)
        _temp_path = data[pos:end]
        pos = end + 1
        size = struct.unpack_from("<I", data, pos)[0]
        pos += 4
        payload = data[pos:pos + size]
        if not payload:
            return None
        return (safe_name(label, "ole_package.bin"), payload)
    except Exception:
        return None


def extract_ole_container(path: Path, output_dir: Path) -> list[Path]:
    try:
        import olefile
    except ImportError:
        return []
    extracted: list[Path] = []
    if not olefile.isOleFile(str(path)):
        return extracted
    with olefile.OleFileIO(str(path)) as ole:
        streams = ole.listdir(streams=True, storages=False)
        for stream_parts in streams:
            stream_name = "/".join(stream_parts)
            lower = stream_name.lower()
            if lower.endswith("ole10native"):
                parsed = parse_ole10native(ole.openstream(stream_parts).read())
                if parsed:
                    name, payload = parsed
                    extracted.append(write_bytes_validated(output_dir, name, payload))
            elif lower.endswith("package") or lower.endswith("contents"):
                payload = ole.openstream(stream_parts).read()
                if payload:
                    extracted.append(write_bytes_validated(output_dir, Path(stream_name).name + ".bin", payload))
    return extracted


def extract_external_links_from_ooxml(path: Path) -> list[str]:
    """Read only explicit external hyperlink relationships from OOXML."""
    links: list[str] = []
    with zipfile.ZipFile(path) as zf:
        for member in zf.namelist():
            if not member.endswith(".rels"):
                continue
            try:
                root = ET.fromstring(zf.read(member))
                for rel in root:
                    rel_type = rel.attrib.get("Type", "")
                    target = rel.attrib.get("Target", "").strip()
                    if (rel.attrib.get("TargetMode") == "External"
                            and rel_type.endswith("/hyperlink")
                            and target.startswith(("http://", "https://"))):
                        links.append(target)
            except Exception:
                continue
    return list(dict.fromkeys(links))


def filename_from_headers_or_url(headers: Any, url: str) -> str:
    disposition = headers.get("Content-Disposition", "") if headers else ""
    match = re.search(
        r"filename\*=UTF-8''([^;]+)|filename=\"?([^\";]+)",
        disposition,
        flags=re.I,
    )
    if match:
        return safe_name(urllib.parse.unquote(match.group(1) or match.group(2)))
    name = Path(urllib.parse.urlparse(url).path).name
    return safe_name(name or "downloaded_document.bin")


def read_limited_response(response: Any) -> bytes:
    length = response.headers.get("Content-Length")
    if length and int(length) > MAX_DOWNLOAD_BYTES:
        raise RuntimeError(f"Linked file is too large: {length} bytes")
    chunks: list[bytes] = []
    total = 0
    while True:
        chunk = response.read(1024 * 1024)
        if not chunk:
            break
        total += len(chunk)
        if total > MAX_DOWNLOAD_BYTES:
            raise RuntimeError("Linked file exceeded MAX_DOWNLOAD_BYTES")
        chunks.append(chunk)
    return b"".join(chunks)


def looks_like_login_page(data: bytes, content_type: str) -> bool:
    if "text/html" not in content_type.lower():
        return False
    sample = data[:200000].decode("utf-8", errors="ignore").lower()
    markers = (
        "login.microsoftonline.com", "sign in to your account", "microsoft login",
        "access denied", "request access", "you need permission", "sharepoint error",
    )
    return any(marker in sample for marker in markers)


def sharepoint_graph_download(url: str, output_dir: Path,
                              token: str) -> tuple[Path | None, str, str | None]:
    """Download a protected SharePoint sharing URL through Microsoft Graph."""
    encoded = base64.urlsafe_b64encode(url.encode("utf-8")).decode("ascii").rstrip("=")
    share_id = "u!" + encoded
    metadata_url = (
        "https://graph.microsoft.com/v1.0/shares/"
        + share_id
        + "/driveItem?$select=id,name,file,folder,@microsoft.graph.downloadUrl"
    )
    request = urllib.request.Request(
        metadata_url,
        headers={"Authorization": f"Bearer {token}", "Accept": "application/json"},
    )
    try:
        with urllib.request.urlopen(request, timeout=DOWNLOAD_TIMEOUT_SECONDS) as response:
            metadata = json.loads(read_limited_response(response).decode("utf-8"))
    except urllib.error.HTTPError as exc:
        if exc.code in {401, 403}:
            return None, "permission_denied", f"SharePoint/Graph returned HTTP {exc.code}"
        return None, "download_failed", f"Microsoft Graph returned HTTP {exc.code}"
    except Exception as exc:
        return None, "download_failed", str(exc)

    if metadata.get("folder") is not None:
        return None, "unsupported_folder", "The SharePoint link points to a folder, not a file"
    download_url = metadata.get("@microsoft.graph.downloadUrl")
    if not download_url:
        return None, "permission_denied", "No authorized download URL was returned"

    try:
        req = urllib.request.Request(download_url, headers={"User-Agent": "DocumentExtractor/2.0"})
        with urllib.request.urlopen(req, timeout=DOWNLOAD_TIMEOUT_SECONDS) as response:
            data = read_limited_response(response)
            content_type = response.headers.get("Content-Type", "")
        name = safe_name(metadata.get("name") or "sharepoint_document.bin")
        saved = write_bytes_validated(output_dir, name, data, content_type)
        return saved, "downloaded", None
    except urllib.error.HTTPError as exc:
        status = "permission_denied" if exc.code in {401, 403} else "download_failed"
        return None, status, f"SharePoint content returned HTTP {exc.code}"
    except Exception as exc:
        return None, "download_failed", str(exc)


def download_external_document(url: str, output_dir: Path
                               ) -> tuple[Path | None, str, str | None]:
    """Download a public link, or use Graph for protected SharePoint links."""
    host = urllib.parse.urlparse(url).netloc.lower()
    is_sharepoint = "sharepoint.com" in host or "1drv.ms" in host

    # A token is the reliable route for protected SharePoint links.
    if is_sharepoint and SHAREPOINT_ACCESS_TOKEN.strip():
        return sharepoint_graph_download(
            url, output_dir, SHAREPOINT_ACCESS_TOKEN.strip()
        )

    request = urllib.request.Request(
        url,
        headers={
            "User-Agent": "Mozilla/5.0 DocumentExtractor/2.0",
            "Accept": "*/*",
        },
    )
    try:
        with urllib.request.urlopen(request, timeout=DOWNLOAD_TIMEOUT_SECONDS) as response:
            final_url = response.geturl()
            content_type = response.headers.get("Content-Type", "")
            data = read_limited_response(response)
            name = filename_from_headers_or_url(response.headers, final_url)
    except urllib.error.HTTPError as exc:
        if exc.code in {401, 403}:
            label = "SharePoint permission/authentication issue" if is_sharepoint else "Permission issue"
            return None, "permission_denied", f"{label}: HTTP {exc.code}"
        return None, "download_failed", f"HTTP {exc.code}: {exc.reason}"
    except urllib.error.URLError as exc:
        return None, "download_failed", f"Network error: {exc.reason}"
    except Exception as exc:
        return None, "download_failed", str(exc)

    if looks_like_login_page(data, content_type):
        message = (
            "SharePoint authentication or file permission is required. "
            "Provide SHAREPOINT_ACCESS_TOKEN with permission to read the file."
            if is_sharepoint
            else "The link redirected to a login or access-denied page."
        )
        return None, "permission_denied", message

    # Do not save a normal web page as though it were an embedded document.
    if "text/html" in content_type.lower():
        return None, "web_page_not_document", (
            "The URL returned an HTML web page instead of a downloadable document"
        )

    saved = write_bytes_validated(output_dir, name, data, content_type)
    return saved, "downloaded", None


def extract_ooxml_embedded(path: Path, document_output_dir: Path,
                           image_output_dir: Path) -> list[Path]:
    """Extract embedded packages and media from DOCX/XLSX/PPTX packages."""
    extracted: list[Path] = []
    with zipfile.ZipFile(path) as zf:
        for member in zf.namelist():
            normalized = member.lower()
            if member.endswith("/"):
                continue

            if "/embeddings/" in normalized or normalized.startswith("embeddings/"):
                data = zf.read(member)
                saved = write_bytes_validated(
                    document_output_dir, Path(member).name, data
                )
                if (saved.suffix.lower() in {".bin", ".ole"}
                        or read_head(saved, 8) == b"\xD0\xCF\x11\xE0\xA1\xB1\x1A\xE1"):
                    unpacked = extract_ole_container(saved, document_output_dir)
                    extracted.extend(unpacked or [saved])
                else:
                    extracted.append(saved)

            # Word, Excel, and PowerPoint store inserted pictures under media/.
            elif "/media/" in normalized or normalized.startswith("media/"):
                data = zf.read(member)
                saved = write_bytes_validated(
                    image_output_dir, Path(member).name, data
                )
                extracted.append(saved)
    return extracted


def extract_pdf_embedded(path: Path, document_output_dir: Path,
                         image_output_dir: Path) -> list[Path]:
    """Extract PDF attachments and unique embedded raster images."""
    import pymupdf
    extracted: list[Path] = []
    seen_xrefs: set[int] = set()
    doc = pymupdf.open(path)
    try:
        # Document-level EmbeddedFiles name tree.
        for i in range(doc.embfile_count()):
            info = doc.embfile_info(i)
            key = info.get("name") or i
            data = doc.embfile_get(key)
            filename = (info.get("filename") or info.get("ufilename")
                        or info.get("name") or f"attachment_{i + 1}.bin")
            extracted.append(
                write_bytes_validated(document_output_dir, filename, data)
            )

        for page in doc:
            # Page-level file attachment annotations.
            annot = page.first_annot
            while annot:
                try:
                    if annot.type[0] == pymupdf.PDF_ANNOT_FILE_ATTACHMENT:
                        info = annot.file_info
                        data = annot.get_file()
                        filename = (info.get("filename") or info.get("ufilename")
                                    or f"page_{page.number + 1}_attachment.bin")
                        extracted.append(
                            write_bytes_validated(document_output_dir, filename, data)
                        )
                except Exception:
                    pass
                annot = annot.next

            # Extract each underlying image object once even if used repeatedly.
            for image_no, image_info in enumerate(page.get_images(full=True), 1):
                xref = image_info[0]
                if xref in seen_xrefs:
                    continue
                seen_xrefs.add(xref)
                try:
                    image = doc.extract_image(xref)
                    ext = image.get("ext", "png")
                    name = f"page_{page.number + 1}_image_{image_no}_xref_{xref}.{ext}"
                    extracted.append(
                        write_bytes_validated(
                            image_output_dir, name, image["image"],
                            f"image/{ext}"
                        )
                    )
                except Exception:
                    # Some PDF image objects are masks or unsupported streams.
                    continue
    finally:
        doc.close()
    return extracted


def process_document(path: Path, root_output: Path, items: list[ExtractedItem],
                     seen_hashes: set[str], depth: int, source: str) -> None:
    """Extract and parse one file, then recursively process every child file."""
    if depth > MAX_RECURSION_DEPTH:
        items.append(ExtractedItem(
            source, str(path), "limit",
            error="Maximum recursion depth reached",
        ))
        return

    path = ensure_correct_extension(path)

    try:
        digest = file_hash(path)
    except Exception as exc:
        items.append(ExtractedItem(
            source, str(path), "file",
            error=f"Hashing failed: {exc}",
        ))
        return

    if digest in seen_hashes:
        return
    seen_hashes.add(digest)

    children: list[Path] = []
    external_links: list[str] = []
    document_dir = root_output / "embedded_files" / f"level_{depth}"
    image_dir = root_output / "extracted_images" / f"level_{depth}"

    # 1. Extract embedded files and images from this container.
    try:
        if path.suffix.lower() == ".pdf":
            children.extend(
                extract_pdf_embedded(path, document_dir, image_dir)
            )
        elif zipfile.is_zipfile(path) and detect_ooxml_extension(path):
            children.extend(
                extract_ooxml_embedded(path, document_dir, image_dir)
            )
            external_links = extract_external_links_from_ooxml(path)
    except Exception as exc:
        items.append(ExtractedItem(
            source, str(path), "container_error", error=str(exc)
        ))

    # 2. Parse this file's own content.
    try:
        text = extract_text(path, root_output / "_work")
        kind = "image_ocr" if path.suffix.lower() in {
            ".png", ".jpg", ".jpeg", ".bmp",
            ".tif", ".tiff", ".gif", ".webp",
        } else "document"
        items.append(ExtractedItem(
            source=source,
            saved_path=str(path),
            kind=kind,
            text=text or "",
        ))
    except Exception as exc:
        items.append(ExtractedItem(
            source, str(path), "document", error=str(exc)
        ))

    # 3. Download every external hyperlink that returns a real file.
    linked_dir = root_output / "external_link_files" / f"level_{depth}"
    for link in external_links:
        if not DOWNLOAD_EXTERNAL_LINKS:
            items.append(ExtractedItem(
                source=str(path),
                saved_path=None,
                kind="external_link",
                link=link,
                access_status="not_downloaded",
            ))
            continue

        downloaded, status, message = download_external_document(
            link, linked_dir
        )
        items.append(ExtractedItem(
            source=str(path),
            saved_path=str(downloaded) if downloaded else None,
            kind="external_link",
            error=message,
            link=link,
            access_status=status,
        ))

        # Critical: downloaded link files enter the same recursive parser as
        # local/embedded files. Their text therefore reaches ALL_CONTENT_TEXT.
        if downloaded and downloaded.is_file():
            children.append(downloaded)

    # 4. Recursively parse embedded files, extracted images and linked files.
    for child in children:
        process_document(
            child,
            root_output,
            items,
            seen_hashes,
            depth + 1,
            str(path),
        )

def main() -> dict[str, Any]:
    global EXTRACTED_CONTENTS, ALL_CONTENT_TEXT, RESULT

    input_path = Path(INPUT_PATH).expanduser().resolve()
    if not input_path.is_file():
        raise FileNotFoundError(
            f"INPUT_PATH does not exist or is not a file: {input_path}\n"
            "Edit INPUT_PATH near the top of this Python file."
        )

    output_root = Path(OUTPUT_FOLDER_NAME)
    if not output_root.is_absolute():
        output_root = input_path.parent / output_root
    output_root = output_root.resolve()

    if output_root.exists():
        shutil.rmtree(output_root)
    output_root.mkdir(parents=True)

    # Keep an unchanged copy of the original with its original name.
    original_copy_dir = output_root / "original_document"
    original_copy_dir.mkdir()
    original_copy = original_copy_dir / input_path.name
    shutil.copy2(input_path, original_copy)

    items: list[ExtractedItem] = []
    process_document(original_copy, output_root, items, set(), 0, "input")

    # Build a stable combined text collection from every successfully parsed
    # original, embedded, OCR-extracted and externally downloaded file.
    EXTRACTED_CONTENTS = {}
    for item in items:
        if item.kind not in {"document", "image_ocr"}:
            continue
        if not item.saved_path or not item.text or not item.text.strip():
            continue

        clean_text = item.text.strip()
        if item.saved_path not in EXTRACTED_CONTENTS:
            EXTRACTED_CONTENTS[item.saved_path] = clean_text
        elif clean_text not in EXTRACTED_CONTENTS[item.saved_path]:
            EXTRACTED_CONTENTS[item.saved_path] += "\n\n" + clean_text

    text_sections: list[str] = []
    for saved_path, text in EXTRACTED_CONTENTS.items():
        text_sections.append(
            "\n\n"
            + "=" * 80
            + f"\nFILE: {saved_path}\n"
            + "=" * 80
            + f"\n{text}"
        )

    # Include link status after content. This makes missing/blocked links clear
    # without replacing content from links that downloaded successfully.
    link_items = [item for item in items if item.kind == "external_link"]
    if link_items:
        text_sections.append(
            "\n\n" + "=" * 80
            + "\nEXTERNAL LINK PROCESSING STATUS\n"
            + "=" * 80
        )
        for item in link_items:
            text_sections.append(
                f"\n\nLink: {item.link or ''}"
                f"\nStatus: {item.access_status or 'unknown'}"
                f"\nSaved file: {item.saved_path or 'Not downloaded'}"
                f"\nMessage: {item.error or 'No error reported'}"
            )

    ALL_CONTENT_TEXT = "".join(text_sections).strip()

    errors = [asdict(item) for item in items if item.error]
    external_link_results = [
        asdict(item) for item in items if item.kind == "external_link"
    ]
    RESULT = {
        "input_file": str(input_path),
        "output_folder": str(output_root),
        "extracted_contents": EXTRACTED_CONTENTS,
        "combined_text": ALL_CONTENT_TEXT,
        "external_links": external_link_results,
        "items": [asdict(item) for item in items],
        "errors": errors,
    }

    # Human-readable text and machine-readable result.
    (output_root / "all_extracted_content.txt").write_text(ALL_CONTENT_TEXT, encoding="utf-8")
    (output_root / "result.json").write_text(
        json.dumps(RESULT, ensure_ascii=False, indent=2, default=str),
        encoding="utf-8",
    )
    (output_root / "external_link_status.json").write_text(
        json.dumps(external_link_results, ensure_ascii=False, indent=2, default=str),
        encoding="utf-8",
    )

    print(f"Completed. Output folder: {output_root}")
    print(f"Documents with extracted text: {len(EXTRACTED_CONTENTS)}")
    print(f"External links found: {len(external_link_results)}")
    for item in external_link_results:
        print(f"  - {item['access_status']}: {item['link']}")
        if item.get("error"):
            print(f"    {item['error']}")
    print(f"Errors/warnings: {len(errors)}")
    if PRINT_CONTENT_TO_CONSOLE:
        print("\n" + "=" * 30 + " EXTRACTED CONTENT " + "=" * 30)
        print(ALL_CONTENT_TEXT or "[No readable text was extracted]")

    return RESULT


if __name__ == "__main__":
    main()
